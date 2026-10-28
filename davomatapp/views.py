import calendar
import csv
from datetime import date, datetime

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import GroupForm, ImportForm, StudentEditForm, StudentForm
from .models import Attendance, Group, Student

OYLAR = [
    "", "Yanvar", "Fevral", "Mart", "Aprel", "May", "Iyun",
    "Iyul", "Avgust", "Sentabr", "Oktabr", "Noyabr", "Dekabr",
]

# Fayldagi sarlavha qatorini avtomatik o'tkazib yuborish uchun
HEADER_WORDS = {
    "f.i.sh", "fish", "f.i.o", "fio", "ism", "ismi", "familiya", "ism familiya",
    "familiya ism", "ism-familiya", "name", "full name", "talaba", "o'quvchi",
    "oquvchi", "o‘quvchi", "t/r", "№",
}


def _parse_date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return timezone.localdate()


def _parse_month(value):
    try:
        d = datetime.strptime(value, "%Y-%m")
        return d.year, d.month
    except (TypeError, ValueError):
        today = timezone.localdate()
        return today.year, today.month


def _shift_month(year, month, delta):
    index = year * 12 + (month - 1) + delta
    return index // 12, index % 12 + 1


# ---------------------------------------------------------------- Bosh sahifa
@login_required
def group_list(request):
    """Guruhlar ro'yxati + bugungi kun statistikasi."""
    today = timezone.localdate()
    groups = list(
        Group.objects.annotate(
            student_count=Count("students", filter=Q(students__is_active=True))
        )
    )

    stats = {
        row["student__group_id"]: row
        for row in Attendance.objects.filter(date=today, student__is_active=True)
        .values("student__group_id")
        .annotate(
            present=Count("id", filter=Q(is_present=True)),
            absent=Count("id", filter=Q(is_present=False)),
        )
    }

    absent_names = {}
    absent_records = Attendance.objects.filter(
        date=today, is_present=False, student__is_active=True
    ).select_related("student")
    for rec in absent_records:
        absent_names.setdefault(rec.student.group_id, []).append(rec.student.full_name)

    total = {"students": 0, "present": 0, "absent": 0, "unmarked": 0}
    for g in groups:
        s = stats.get(g.id, {})
        g.present_today = s.get("present", 0)
        g.absent_today = s.get("absent", 0)
        marked = g.present_today + g.absent_today
        g.marked = marked > 0
        g.unmarked_today = max(g.student_count - marked, 0)
        g.percent_today = round(g.present_today * 100 / marked) if marked else 0
        g.absent_names = absent_names.get(g.id, [])

        total["students"] += g.student_count
        total["present"] += g.present_today
        total["absent"] += g.absent_today
        total["unmarked"] += g.unmarked_today

    marked_total = total["present"] + total["absent"]
    total["percent"] = round(total["present"] * 100 / marked_total) if marked_total else None

    return render(
        request,
        "davomatapp/group_list.html",
        {"groups": groups, "total": total, "today": today},
    )


# ------------------------------------------------------------------- Guruhlar
@login_required
def group_create(request):
    form = GroupForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        group = form.save()
        messages.success(request, f"“{group.name}” guruhi yaratildi.")
        return redirect("group_detail", pk=group.pk)
    return render(request, "davomatapp/form.html", {"form": form, "title": "Yangi guruh"})


@login_required
def group_detail(request, pk):
    """Oylik davomat jadvali: kunlar ustunda, talabalar qatorda."""
    group = get_object_or_404(Group, pk=pk)
    year, month = _parse_month(request.GET.get("oy"))
    days_in_month = calendar.monthrange(year, month)[1]
    days = [date(year, month, d) for d in range(1, days_in_month + 1)]

    students = list(group.students.filter(is_active=True))
    records = Attendance.objects.filter(
        student__in=students, date__year=year, date__month=month
    )
    table = {(r.student_id, r.date): r.is_present for r in records}
    lesson_days = {d for (_, d) in table.keys()}

    rows = []
    for s in students:
        cells = []
        present = absent = 0
        for d in days:
            value = table.get((s.id, d))
            cells.append(value)
            if value is True:
                present += 1
            elif value is False:
                absent += 1
        total = present + absent
        rows.append(
            {
                "student": s,
                "cells": cells,
                "present": present,
                "absent": absent,
                "percent": round(present * 100 / total) if total else None,
            }
        )

    prev_y, prev_m = _shift_month(year, month, -1)
    next_y, next_m = _shift_month(year, month, 1)
    context = {
        "group": group,
        "days": days,
        "rows": rows,
        "lesson_days_count": len(lesson_days),
        "month_title": f"{OYLAR[month]} {year}",
        "month_value": f"{year}-{month:02d}",
        "prev_month": f"{prev_y}-{prev_m:02d}",
        "next_month": f"{next_y}-{next_m:02d}",
        "today": timezone.localdate(),
    }
    return render(request, "davomatapp/group_detail.html", context)


# ------------------------------------------------------------------ Talabalar
@login_required
def student_create(request, pk):
    group = get_object_or_404(Group, pk=pk)
    form = StudentForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        student = form.save(commit=False)
        student.group = group
        student.save()
        messages.success(request, f"{student.full_name} guruhga qo'shildi.")
        return redirect("group_detail", pk=group.pk)
    return render(
        request,
        "davomatapp/form.html",
        {"form": form, "title": f"{group.name}: yangi talaba", "back": group},
    )


@login_required
def student_update(request, pk):
    student = get_object_or_404(Student, pk=pk)
    form = StudentEditForm(request.POST or None, instance=student)
    if request.method == "POST" and form.is_valid():
        student = form.save()
        messages.success(request, f"{student.full_name} ma'lumotlari yangilandi.")
        return redirect("group_detail", pk=student.group_id)
    return render(
        request,
        "davomatapp/form.html",
        {
            "form": form,
            "title": f"Tahrirlash: {student.full_name}",
            "back": student.group,
            "delete_url": reverse("student_delete", args=[student.pk]),
        },
    )


@login_required
@require_POST
def student_delete(request, pk):
    student = get_object_or_404(Student, pk=pk)
    group_id = student.group_id
    name = student.full_name
    student.delete()
    messages.success(request, f"{name} o'chirildi.")
    return redirect("group_detail", pk=group_id)


@login_required
def student_detail(request, pk):
    student = get_object_or_404(Student, pk=pk)
    records = student.attendances.all()
    present = records.filter(is_present=True).count()
    absent = records.filter(is_present=False).count()
    total = present + absent
    return render(
        request,
        "davomatapp/student_detail.html",
        {
            "student": student,
            "records": records,
            "present": present,
            "absent": absent,
            "percent": round(present * 100 / total) if total else None,
        },
    )


@login_required
def student_stats(request, pk):
    """Ism ustiga bosilganda oynada chiqadigan statistika (JSON)."""
    student = get_object_or_404(Student.objects.select_related("group"), pk=pk)
    records = student.attendances.all()
    present = records.filter(is_present=True).count()
    absent = records.filter(is_present=False).count()
    total = present + absent

    today = timezone.localdate()
    month_qs = records.filter(date__year=today.year, date__month=today.month)
    m_present = month_qs.filter(is_present=True).count()
    m_absent = month_qs.filter(is_present=False).count()
    m_total = m_present + m_absent

    recent = [
        {"date": r.date.strftime("%d.%m.%Y"), "present": r.is_present, "note": r.note}
        for r in records[:8]
    ]
    return JsonResponse(
        {
            "name": student.full_name,
            "group": student.group.name,
            "phone": student.phone,
            "present": present,
            "absent": absent,
            "percent": round(present * 100 / total) if total else None,
            "month_title": f"{OYLAR[today.month]} {today.year}",
            "month_present": m_present,
            "month_absent": m_absent,
            "month_percent": round(m_present * 100 / m_total) if m_total else None,
            "recent": recent,
            "detail_url": reverse("student_detail", args=[student.pk]),
            "edit_url": reverse("student_update", args=[student.pk]),
        }
    )


# ------------------------------------------------------------ Fayldan yuklash
def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_rows(uploaded):
    """Fayldan qatorlar ro'yxatini o'qiydi: [[ustun1, ustun2, ...], ...]"""
    name = uploaded.name.lower()

    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError("Excel o'qish uchun terminalda bajaring: pip install openpyxl")
        try:
            wb = load_workbook(uploaded, read_only=True, data_only=True)
            ws = wb.active
            rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
            wb.close()
            return rows
        except Exception:
            raise ValueError("Excel faylini o'qib bo'lmadi. Fayl buzilmaganini tekshiring.")

    if name.endswith((".csv", ".txt")):
        raw = uploaded.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1251", errors="replace")
        lines = [ln for ln in text.splitlines() if ln.strip()]
        if not lines:
            return []
        delimiter = max([";", "\t", ","], key=lines[0].count)
        return [[c.strip() for c in row] for row in csv.reader(lines, delimiter=delimiter)]

    raise ValueError("Faqat .xlsx, .csv yoki .txt fayl yuklang.")


def _import_students(group, rows):
    existing = {n.lower() for n in group.students.values_list("full_name", flat=True)}
    new_students, skipped = [], 0

    for cells in rows:
        # 1-ustun tartib raqami bo'lsa (1, 2, 3...), uni tashlab yuboramiz
        if len(cells) > 1 and cells[0].isdigit():
            cells = cells[1:]
        if not cells or not cells[0]:
            continue

        name = " ".join(cells[0].split())
        if name.isdigit() or name.lower().strip(" .:") in HEADER_WORDS:
            continue

        key = name.lower()
        if key in existing:
            skipped += 1
            continue
        existing.add(key)

        phone = cells[1][:20] if len(cells) > 1 else ""
        new_students.append(Student(full_name=name[:150], group=group, phone=phone))

    Student.objects.bulk_create(new_students)
    return len(new_students), skipped


@login_required
def group_import(request, pk):
    group = get_object_or_404(Group, pk=pk)
    form = ImportForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            rows = _read_rows(form.cleaned_data["file"])
        except ValueError as e:
            form.add_error("file", str(e))
        else:
            created, skipped = _import_students(group, rows)
            if created:
                msg = f"{created} ta talaba qo'shildi."
                if skipped:
                    msg += f" {skipped} tasi guruhda bor edi, o'tkazib yuborildi."
                messages.success(request, msg)
            else:
                messages.success(request, "Yangi talaba topilmadi (hammasi guruhda bor yoki fayl bo'sh).")
            return redirect("group_detail", pk=group.pk)
    return render(request, "davomatapp/import.html", {"form": form, "group": group})


# ------------------------------------------------------------------- Davomat
@login_required
def mark_attendance(request, pk):
    """Bir kunlik davomatni belgilash."""
    group = get_object_or_404(Group, pk=pk)
    students = list(group.students.filter(is_active=True))

    if request.method == "POST":
        lesson_date = _parse_date(request.POST.get("date"))
        for s in students:
            Attendance.objects.update_or_create(
                student=s,
                date=lesson_date,
                defaults={
                    "is_present": request.POST.get(f"present_{s.id}") == "on",
                    "note": request.POST.get(f"note_{s.id}", "").strip(),
                },
            )
        messages.success(request, f"{lesson_date:%d.%m.%Y} sanasi uchun davomat saqlandi.")
        url = reverse("group_detail", args=[group.pk])
        return redirect(f"{url}?oy={lesson_date:%Y-%m}")

    lesson_date = _parse_date(request.GET.get("date"))
    existing = {
        a.student_id: a
        for a in Attendance.objects.filter(student__in=students, date=lesson_date)
    }
    items = []
    for s in students:
        rec = existing.get(s.id)
        items.append(
            {
                "student": s,
                "present": rec.is_present if rec else True,
                "note": rec.note if rec else "",
            }
        )
    return render(
        request,
        "davomatapp/mark.html",
        {
            "group": group,
            "items": items,
            "lesson_date": lesson_date,
            "already_saved": bool(existing),
        },
    )


@login_required
def export_csv(request, pk):
    group = get_object_or_404(Group, pk=pk)
    year, month = _parse_month(request.GET.get("oy"))
    days_in_month = calendar.monthrange(year, month)[1]

    students = list(group.students.filter(is_active=True))
    records = Attendance.objects.filter(
        student__in=students, date__year=year, date__month=month
    )
    table = {(r.student_id, r.date.day): r.is_present for r in records}

    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = (
        f'attachment; filename="davomat_{group.pk}_{year}-{month:02d}.csv"'
    )
    response.write("\ufeff")  # Excel uchun UTF-8 BOM
    writer = csv.writer(response)
    writer.writerow(["F.I.Sh."] + list(range(1, days_in_month + 1)) + ["Keldi", "Kelmadi"])
    for s in students:
        row, p, a = [], 0, 0
        for d in range(1, days_in_month + 1):
            v = table.get((s.id, d))
            if v is True:
                row.append("+")
                p += 1
            elif v is False:
                row.append("-")
                a += 1
            else:
                row.append("")
        writer.writerow([s.full_name] + row + [p, a])
    return response